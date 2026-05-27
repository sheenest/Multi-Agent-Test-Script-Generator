"""
Generator Agent — Takes the Analyzer's test plan and produces the final
Excel test case file in the exact format matching the existing CR test files.

Output format matches the uploaded CR test case Excel files:
  Sheet 1: "Test Cases by Component" — grouped test cases with all columns
  Sheet 2: "Traceability Matrix" — component to test case mapping
  Sheet 3: "Summary" — counts by type

Tool: save_test_excel — builds and saves the formatted .xlsx file
"""

import json
import re
from pathlib import Path
from strands import Agent, tool
from strands.models import BedrockModel
# from shared import LLM_MODEL, BEDROCK_REGION

import sys
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parent.parent))

from shared import embed, get_index, LLM_MODEL, BEDROCK_REGION , get_model

@tool
def save_test_excel(cr_id: str, cr_title: str, test_plan_json: str) -> str:
    """Save the generated test cases as a formatted Excel file.

    The test_plan_json must be a JSON string with this exact structure:
    {
        "groups": [
            {
                "group_num": "1",
                "component": "Expense Item Configuration",
                "design_ref": "3.1",
                "test_cases": [
                    {
                        "tc_num": "P-01",
                        "test_type": "Positive",
                        "scenario": "Item visible in picker",
                        "steps": "User opens new expense report...",
                        "expected_result": "Item appears under group..."
                    }
                ]
            }
        ]
    }

    Args:
        cr_id: The CR identifier (e.g., 'CR-EXP-2026-NEW').
        cr_title: The descriptive title of the CR.
        test_plan_json: JSON string containing the structured test groups.
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    try:
        plan = json.loads(test_plan_json)
    except json.JSONDecodeError as e:
        return f"ERROR: Invalid JSON \u2014 {e}"

    groups = plan.get("groups", [])
    if not groups:
        return "ERROR: No test groups found in the plan."

    wb = openpyxl.Workbook()

    # ── Styles ────────────────────────────────────────────────────────
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(
        start_color="4472C4", end_color="4472C4", fill_type="solid"
    )
    header_font_white = Font(bold=True, size=11, color="FFFFFF")
    title_font = Font(bold=True, size=14)
    subtitle_font = Font(bold=True, size=12, color="4472C4")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    wrap = Alignment(wrap_text=True, vertical="top")

    # === Sheet 1: Test Cases by Component =============================
    ws1 = wb.active
    ws1.title = "Test Cases by Component"

    # Title rows
    ws1.cell(row=1, column=1,
             value=f"{cr_id} \u2014 Test Cases").font = title_font
    ws1.cell(row=2, column=1, value=cr_title).font = subtitle_font

    # Header row
    headers = [
        "Group #", "Solution Component", "Design Ref", "TC #",
        "Test Type", "Scenario", "Test Steps", "Expected Result",
    ]
    for col, header in enumerate(headers, 1):
        cell = ws1.cell(row=4, column=col, value=header)
        cell.font = header_font_white
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(
            horizontal="center", vertical="center"
        )

    current_row = 5
    for group in groups:
        for tc in group.get("test_cases", []):
            for col, val in enumerate([
                group["group_num"],
                group["component"],
                group["design_ref"],
                tc["tc_num"],
                tc["test_type"],
                tc["scenario"],
                tc.get("steps", ""),
                tc.get("expected_result", ""),
            ], 1):
                cell = ws1.cell(row=current_row, column=col, value=val)
                cell.border = thin_border
                cell.alignment = wrap
            current_row += 1
        current_row += 1  # blank row between groups

    # Column widths
    col_widths = [8, 40, 12, 8, 10, 35, 50, 45]
    for i, w in enumerate(col_widths, 1):
        ws1.column_dimensions[get_column_letter(i)].width = w

    # === Sheet 2: Traceability Matrix =================================
    ws2 = wb.create_sheet("Traceability Matrix")
    ws2.cell(row=1, column=1,
             value=f"{cr_id} \u2014 Traceability Matrix").font = title_font

    tm_headers = [
        "#", "Solution Component", "Design Ref",
        "Positive", "Negative", "Regression",
    ]
    for col, header in enumerate(tm_headers, 1):
        cell = ws2.cell(row=3, column=col, value=header)
        cell.font = header_font_white
        cell.fill = header_fill
        cell.border = thin_border

    total_pos = total_neg = total_reg = 0
    for row_idx, group in enumerate(groups, 4):
        tcs = group.get("test_cases", [])
        pos = [tc["tc_num"] for tc in tcs
               if tc["test_type"] == "Positive"]
        neg = [tc["tc_num"] for tc in tcs
               if tc["test_type"] == "Negative"]
        reg = [tc["tc_num"] for tc in tcs
               if tc["test_type"] == "Regression"]
        total_pos += len(pos)
        total_neg += len(neg)
        total_reg += len(reg)

        for col, val in enumerate([
            group["group_num"],
            group["component"],
            group["design_ref"],
            ", ".join(pos) or "\u2014",
            ", ".join(neg) or "\u2014",
            ", ".join(reg) or "\u2014",
        ], 1):
            cell = ws2.cell(row=row_idx, column=col, value=val)
            cell.border = thin_border
            cell.alignment = wrap

    # Total row
    total_row = 4 + len(groups)
    ws2.cell(row=total_row, column=2, value="TOTAL").font = header_font
    ws2.cell(row=total_row, column=4, value=total_pos).font = header_font
    ws2.cell(row=total_row, column=5, value=total_neg).font = header_font
    ws2.cell(row=total_row, column=6, value=total_reg).font = header_font
    for col in range(1, 7):
        ws2.cell(row=total_row, column=col).border = thin_border

    tm_widths = [6, 45, 12, 20, 20, 25]
    for i, w in enumerate(tm_widths, 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    # === Sheet 3: Summary =============================================
    ws3 = wb.create_sheet("Summary")
    ws3.cell(row=1, column=1,
             value=f"{cr_id} \u2014 Test Summary").font = title_font

    total_all = total_pos + total_neg + total_reg
    summary_data = [
        ("Total Test Cases", total_all),
        ("Positive Tests", total_pos),
        ("Negative Tests", total_neg),
        ("Regression Tests", total_reg),
        ("", ""),
        ("Solution Components", len(groups)),
        ("Groups with Positive", sum(
            1 for g in groups if any(
                tc["test_type"] == "Positive"
                for tc in g.get("test_cases", [])
            )
        )),
        ("Groups with Negative", sum(
            1 for g in groups if any(
                tc["test_type"] == "Negative"
                for tc in g.get("test_cases", [])
            )
        )),
    ]

    ws3.cell(row=3, column=1, value="Metric").font = header_font_white
    ws3.cell(row=3, column=1).fill = header_fill
    ws3.cell(row=3, column=2, value="Count").font = header_font_white
    ws3.cell(row=3, column=2).fill = header_fill

    for i, (metric, count) in enumerate(summary_data, 4):
        ws3.cell(row=i, column=1, value=metric).border = thin_border
        ws3.cell(row=i, column=2, value=count).border = thin_border

    ws3.column_dimensions["A"].width = 25
    ws3.column_dimensions["B"].width = 12

    # ── Save ──────────────────────────────────────────────────────────
    out_dir = Path("generated_tests")
    out_dir.mkdir(exist_ok=True)
    filename = f"{cr_id}_TestCases.xlsx"
    filepath = out_dir / filename
    wb.save(str(filepath))

    return (
        f"\u2713 Saved {filepath}\n"
        f"  Total test cases: {total_all} "
        f"({total_pos}P / {total_neg}N / {total_reg}R)\n"
        f"  Groups: {len(groups)}\n"
        f"  Sheets: Test Cases by Component, Traceability Matrix, Summary"
    )


# ── Agent Definition ──────────────────────────────────────────────────
generator_agent = Agent(
    name="generator",
    # model=BedrockModel(model_id=LLM_MODEL, region_name=BEDROCK_REGION),
    model = get_model(),
    tools=[save_test_excel],
    system_prompt="""You are a Test Case Generator for Workday Expenses \
Change Requests.

You receive a detailed test plan from the Analyzer agent and reference \
test cases from past CRs. Your job is to produce the FINAL test case \
Excel file.

=== CRITICAL: OUTPUT FORMAT ===

You MUST call save_test_excel with a JSON string in this EXACT structure:

{
    "groups": [
        {
            "group_num": "1",
            "component": "Expense Item Configuration",
            "design_ref": "3.1",
            "test_cases": [
                {
                    "tc_num": "P-01",
                    "test_type": "Positive",
                    "scenario": "Item visible in picker",
                    "steps": "User opens new expense report and browses \
item picker",
                    "expected_result": "Item appears under correct expense \
group"
                }
            ]
        }
    ]
}

=== NUMBERING RULES ===

- Positive tests: P-01, P-02, P-03, etc. (sequential across all groups)
- Negative tests: N-01, N-02, N-03, etc. (sequential across all groups)
- Regression tests: R-01, R-02, R-03, etc. (sequential across all groups)

=== TEST CASE QUALITY RULES ===

For EVERY test case, provide:
- tc_num: Use the P/N/R numbering above
- test_type: Exactly "Positive", "Negative", or "Regression"
- scenario: Short description (5-10 words)
- steps: Detailed Workday-specific steps the tester performs
- expected_result:
  * Positive: "Submits successfully" or "Saves without errors" with specifics
  * Negative: "Hard stop: [exact validation message text]"
  * Regression: "[Existing feature] behaves as before" with specifics

=== WORKDAY EXPENSES CONTEXT ===

- Expense reports contain one or more expense lines
- Each line has: expense item, amount, currency, date, memo, receipt, \
custom fields
- "Submit" triggers all validation rules on the report/lines
- Hard stop = blocks submission with error message
- Soft stop = warning that can be acknowledged to proceed
- Always test in the context of: Create expense report -> add expense \
line -> fill fields -> submit

Call save_test_excel ONCE with the complete test plan. Then provide a \
brief summary.""",
)